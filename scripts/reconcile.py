"""Reconcile two schedule workbooks that disagree into one event list.

    uv run python scripts/reconcile.py "<dir containing both workbooks>"

Neither view is a superset of the other, so neither can be the sole import
source. This prints the merged calendar plus every disagreement it had to
resolve; the resolutions themselves are judgement calls a human must confirm.
"""

import datetime as dt
import sys

import openpyxl

if len(sys.argv) < 2:
    raise SystemExit("usage: reconcile.py <directory containing the workbooks>")
BASE = sys.argv[1].rstrip("/") + "/"
DOW = ['MON', 'TUE', 'WED', 'THU', 'FRI', 'SAT', 'SUN']

# ---------- View A: flat list ----------
wbA = openpyxl.load_workbook(BASE + 'REU-feature-exploration-exhaustive-negative.xlsx', data_only=True)
wsA = wbA['Assignments']
viewA = {}          # date -> text
rawA = []
for row in wsA.iter_rows(min_row=3, values_only=True):
    d, with_, ev, theme, loc = (list(row) + [None] * 5)[:5]
    if not isinstance(d, dt.datetime):
        continue
    txt = ' / '.join(str(x) for x in (with_, ev) if x)
    rawA.append((d.date(), txt))
    viewA[d.date()] = txt

# ---------- View B: month grids ----------
# Layout: a row of day-numbers (ints under B..H = SUN..SAT), then event rows
# beneath it until the next day-number row. Column letter -> weekday offset.
COLS = ['B', 'C', 'D', 'E', 'F', 'G', 'H']
wbB = openpyxl.load_workbook(BASE + 'Full Fall 2026 Schedule.xlsx', data_only=True)
viewB = {}
titles = {}
for ws in wbB.worksheets:
    month = ws.title.split()[0]
    mnum = ['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec'].index(month[:3]) + 1
    titles[ws.title] = ws['B1'].value
    current = {}   # col -> date
    for r in range(2, ws.max_row + 1):
        vals = {c: ws[f'{c}{r}'].value for c in COLS}
        nums = {c: v for c, v in vals.items() if isinstance(v, (int, float)) and 1 <= v <= 31}
        if len(nums) >= 3:                     # this is a day-number row
            current = {c: dt.date(2026, mnum, int(v)) for c, v in nums.items()}
            continue
        for c, v in vals.items():
            if v and c in current and isinstance(v, str):
                viewB.setdefault(current[c], []).append(v.strip())

print('=' * 72)
print('A. SHEET TITLES vs ACTUAL CALENDAR YEAR')
print('=' * 72)
for ws in wbB.worksheets:
    month = ws.title.split()[0]
    mnum = ['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec'].index(month[:3]) + 1
    # find first day-number row to infer the weekday of the 1st
    for r in range(2, 20):
        vals = {c: ws[f'{c}{r}'].value for c in COLS}
        nums = {c: v for c, v in vals.items() if isinstance(v, (int, float)) and 1 <= v <= 31}
        if len(nums) >= 3:
            col1 = [c for c, v in nums.items() if v == 1]
            if col1:
                idx = COLS.index(col1[0])          # 0=SUN
                grid_dow = ['SUN','MON','TUE','WED','THU','FRI','SAT'][idx]
                d25 = dt.date(2025, mnum, 1).strftime('%a').upper()[:3]
                d26 = dt.date(2026, mnum, 1).strftime('%a').upper()[:3]
                print(f'  {ws.title:<10} title={titles[ws.title]!r:<20} '
                      f'grid says 1st = {grid_dow} | 2025={d25} 2026={d26} '
                      f'-> matches {"2026" if grid_dow.startswith(d26) else "2025" if grid_dow.startswith(d25) else "NEITHER"}')
            break

print()
print('=' * 72)
print('B. BAD / OUT-OF-RANGE DATES IN VIEW A')
print('=' * 72)
prev = None
for d, txt in rawA:
    if d.year != 2026:
        print(f'  !! {d} (year {d.year}) — {txt!r}')
    if prev and d < prev:
        print(f'  !! out-of-order: {d} follows {prev}')
    prev = d

print()
print('=' * 72)
print('C. EVENT-BY-EVENT RECONCILIATION (Aug 25 - Dec 5, 2026)')
print('=' * 72)
start, end = dt.date(2026, 8, 25), dt.date(2026, 12, 5)
alldates = sorted(set(list(viewA) + list(viewB)))
IGNORE = ('chapter 7pm',)
for d in alldates:
    if not (start <= d <= end):
        continue
    a = viewA.get(d)
    b = [x for x in viewB.get(d, []) if x.lower() not in IGNORE]
    bt = ' / '.join(b)
    if a and not bt:
        print(f'  {d} {DOW[d.weekday()]}  ONLY IN A : {a!r}')
    elif bt and a is None:
        print(f'  {d} {DOW[d.weekday()]}  ONLY IN B : {bt!r}')
    elif a == '' and bt:
        print(f'  {d} {DOW[d.weekday()]}  A BLANK   : B has {bt!r}')
    elif a and bt:
        print(f'  {d} {DOW[d.weekday()]}  both      : A={a!r} | B={bt!r}')

print()
print('=' * 72)
print('D. EMPTY TUE / FRI / SAT -> PLACEHOLDER CANDIDATES')
print('=' * 72)
d = start
while d <= end:
    if d.weekday() in (1, 4, 5):           # Tue, Fri, Sat
        a = viewA.get(d)
        b = [x for x in viewB.get(d, []) if x.lower() not in IGNORE]
        if not a and not b:
            print(f'  {d} {DOW[d.weekday()]}  — no event in either view')
    d += dt.timedelta(days=1)
